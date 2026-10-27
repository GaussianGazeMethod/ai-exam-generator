"""S6-3 报表导出 — 合规报告 PDF + 统计明细 Excel。

依据 spec/核心模块架构设计.md 模块六 M6-3：
- 合规报告 PDF（spec 要求含本次活动记录：谁 / 时间 / 得分 / 用时 / IP）
- 统计明细 Excel

技术实现（已实测验证）：
- PDF：reportlab + `UnicodeCIDFont('STSong-Light')`
  CID 字体是 reportlab 内置的 CJK 支持，无需外挂字体文件，中文渲染正常
- Excel：openpyxl，多 sheet 分维度导出
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

# ── 中文字体（CID，无需外部字体文件）──
_CJK_FONT = "STSong-Light"
_font_registered = False


def _ensure_font() -> str:
    global _font_registered
    if not _font_registered:
        pdfmetrics.registerFont(UnicodeCIDFont(_CJK_FONT))
        _font_registered = True
    return _CJK_FONT


# ──────────────────────────────────────────────────────────
# PDF：合规报告
# ──────────────────────────────────────────────────────────


def build_compliance_pdf(
    *,
    exam_overview: dict[str, Any],
    dept_stats: list[dict[str, Any]],
    person_stats: list[dict[str, Any]],
    chapter_stats: list[dict[str, Any]] | None = None,
    question_stats: list[dict[str, Any]] | None = None,
    generated_at: datetime | None = None,
) -> bytes:
    """生成合规报告 PDF。

    结构（对应 spec 合规留痕要求）：
    1. 报告头：考试名称、生成时间、考试状态
    2. 概览：应考/已考/未考、平均分、通过率、最高/最低分
    3. 按部门汇总
    4. 合规留痕明细：谁 / 时间 / 得分 / 用时 / IP
    5. 薄弱章节（如有）
    6. 高错误率题目（如有）
    """
    font = _ensure_font()
    generated_at = generated_at or datetime.now()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=f"{exam_overview.get('exam_name', '考试')} 合规报告",
    )

    base = getSampleStyleSheet()
    st_title = ParagraphStyle("t", parent=base["Title"], fontName=font, fontSize=18, spaceAfter=6)
    st_h2 = ParagraphStyle("h2", parent=base["Heading2"], fontName=font, fontSize=13, spaceBefore=14, spaceAfter=8)
    st_body = ParagraphStyle("b", parent=base["Normal"], fontName=font, fontSize=9.5, leading=15)
    st_note = ParagraphStyle("n", parent=st_body, textColor=colors.HexColor("#666666"), fontSize=8.5)

    story: list[Any] = []

    # ── 1. 报告头 ──
    story.append(Paragraph("制度传达验证 — 合规报告", st_title))
    story.append(
        Paragraph(
            f"考试名称：{exam_overview.get('exam_name', '-')}　|　"
            f"考试状态：{exam_overview.get('exam_status', '-')}　|　"
            f"生成时间：{generated_at.strftime('%Y-%m-%d %H:%M:%S')}",
            st_body,
        )
    )
    story.append(
        Paragraph(
            "本报告依据员工线上考试记录生成，用于证明相关制度已完成传达并已验证。",
            st_note,
        )
    )
    story.append(Spacer(1, 8))

    # ── 2. 概览 ──
    story.append(Paragraph("一、总体情况", st_h2))
    overview_rows = [
        ["应考人数", "已考人数", "未考人数", "平均分", "通过率", "最高分", "最低分"],
        [
            str(exam_overview.get("candidate_count", 0)),
            str(exam_overview.get("completed_count", 0)),
            str(exam_overview.get("not_started_count", 0)),
            str(exam_overview.get("average_score", 0)),
            f"{exam_overview.get('pass_rate', 0)}%",
            str(exam_overview.get("highest_score", 0)),
            str(exam_overview.get("lowest_score", 0)),
        ],
    ]
    story.append(_make_table(overview_rows, font, col_widths=[32 * mm] * 7))

    # ── 3. 按部门 ──
    if dept_stats:
        story.append(Paragraph("二、按部门统计", st_h2))
        rows = [["排名", "部门", "应考", "已考", "未考", "平均分", "通过率"]]
        for d in dept_stats:
            rows.append(
                [
                    str(d.get("rank", "-")),
                    str(d.get("dept_name", "-")),
                    str(d.get("candidate_count", 0)),
                    str(d.get("completed_count", 0)),
                    str(d.get("not_started_count", 0)),
                    str(d.get("average_score", 0)),
                    f"{d.get('pass_rate', 0)}%",
                ]
            )
        story.append(_make_table(rows, font, col_widths=[18 * mm, 60 * mm, 24 * mm, 24 * mm, 24 * mm, 26 * mm, 26 * mm]))

    # ── 4. 合规留痕明细（谁 / 时间 / 得分 / 用时 / IP）──
    story.append(PageBreak())
    story.append(Paragraph("三、考试记录明细（合规留痕）", st_h2))
    if person_stats:
        rows = [["姓名", "部门", "得分", "是否通过", "答对/答错", "用时", "提交时间", "IP 地址"]]
        for p in person_stats:
            rows.append(
                [
                    str(p.get("name", "-")),
                    str(p.get("dept_name", "-")),
                    str(p.get("score", 0)),
                    "是" if p.get("is_passed") else "否",
                    f"{p.get('correct_count', 0)} / {p.get('wrong_count', 0)}",
                    _fmt_duration(p.get("duration_seconds") or 0),
                    str(p.get("submitted_at", "")).replace("T", " ")[:19] or "-",
                    str(p.get("ip_address") or "-"),
                ]
            )
        story.append(
            _make_table(
                rows,
                font,
                col_widths=[28 * mm, 34 * mm, 18 * mm, 20 * mm, 24 * mm, 24 * mm, 40 * mm, 34 * mm],
                font_size=8.5,
            )
        )
    else:
        story.append(Paragraph("暂无考生提交记录。", st_body))

    # ── 5. 薄弱章节 ──
    if chapter_stats:
        story.append(Paragraph("四、章节正确率（升序，最薄弱在前）", st_h2))
        rows = [["章节", "节", "作答次数", "正确次数", "正确率"]]
        for c in chapter_stats[:15]:
            rows.append(
                [
                    str(c.get("chapter_no", "-")),
                    str(c.get("section_no") or "-"),
                    str(c.get("total_answers", 0)),
                    str(c.get("correct_answers", 0)),
                    f"{c.get('correct_rate', 0)}%",
                ]
            )
        story.append(_make_table(rows, font, col_widths=[50 * mm, 40 * mm, 30 * mm, 30 * mm, 30 * mm]))

    # ── 6. 高错误率题目 ──
    if question_stats:
        flagged = [q for q in question_stats if q.get("needs_review")]
        if flagged:
            story.append(Paragraph("五、待复查题目（错误率异常，建议核对表述）", st_h2))
            rows = [["题号", "题型", "题目摘要", "错误率", "作答次数"]]
            for q in flagged[:15]:
                rows.append(
                    [
                        str(q.get("order", "-")),
                        _type_label(q.get("question_type", "")),
                        str(q.get("content", ""))[:48],
                        f"{q.get('error_rate', 0)}%",
                        str(q.get("total_answers", 0)),
                    ]
                )
            story.append(
                _make_table(
                    rows,
                    font,
                    col_widths=[16 * mm, 22 * mm, 120 * mm, 24 * mm, 26 * mm],
                    font_size=8.5,
                )
            )

    story.append(Spacer(1, 14))
    story.append(
        Paragraph(
            f"报告生成时间：{generated_at.strftime('%Y-%m-%d %H:%M:%S')}　|　"
            "本报告由 AI 试题生成系统自动生成，记录不可篡改。",
            st_note,
        )
    )

    doc.build(story)
    return buf.getvalue()


def _make_table(
    rows: list[list[str]],
    font: str,
    *,
    col_widths: list[float] | None = None,
    font_size: float = 9,
) -> Table:
    table = Table(rows, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("FONTSIZE", (0, 0), (-1, -1), font_size),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF7")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0D2A4A")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B0BEC5")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (2, 1), (-1, -1), "CENTER"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def _fmt_duration(seconds: int) -> str:
    if not seconds:
        return "-"
    m, s = divmod(int(seconds), 60)
    return f"{m}分{s}秒"


def _type_label(qtype: str) -> str:
    return {
        "single_choice": "单选",
        "multi_choice": "多选",
        "true_false": "判断",
        "fill_blank": "填空",
    }.get(qtype, qtype)


# ──────────────────────────────────────────────────────────
# Excel：统计明细
# ──────────────────────────────────────────────────────────


def build_statistics_excel(
    *,
    exam_overview: dict[str, Any],
    dept_stats: list[dict[str, Any]],
    person_stats: list[dict[str, Any]],
    chapter_stats: list[dict[str, Any]] | None = None,
    question_stats: list[dict[str, Any]] | None = None,
    generated_at: datetime | None = None,
) -> bytes:
    """生成多 sheet 统计明细 Excel。"""
    generated_at = generated_at or datetime.now()
    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def _write_sheet(ws, title: str, headers: list[str], rows: list[list[Any]]) -> None:
        ws.title = title
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        for row in rows:
            ws.append(row)
        # 自适应列宽
        for c in range(1, len(headers) + 1):
            width = len(str(headers[c - 1])) * 2 + 6
            for r in range(2, ws.max_row + 1):
                v = ws.cell(row=r, column=c).value
                if v is not None:
                    width = max(width, min(len(str(v)) + 4, 60))
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.freeze_panes = "A2"

    # Sheet 1：概览
    ws = wb.active
    _write_sheet(
        ws,
        "总体概览",
        ["指标", "数值"],
        [
            ["考试名称", exam_overview.get("exam_name", "-")],
            ["考试状态", exam_overview.get("exam_status", "-")],
            ["应考人数", exam_overview.get("candidate_count", 0)],
            ["已考人数", exam_overview.get("completed_count", 0)],
            ["未考人数", exam_overview.get("not_started_count", 0)],
            ["平均分", exam_overview.get("average_score", 0)],
            ["通过率(%)", exam_overview.get("pass_rate", 0)],
            ["最高分", exam_overview.get("highest_score", 0)],
            ["最低分", exam_overview.get("lowest_score", 0)],
            ["通过人数", exam_overview.get("passed_count", 0)],
            ["未通过人数", exam_overview.get("failed_count", 0)],
            ["导出时间", generated_at.strftime("%Y-%m-%d %H:%M:%S")],
        ],
    )

    # Sheet 2：按部门
    if dept_stats:
        _write_sheet(
            wb.create_sheet(),
            "按部门",
            ["排名", "部门", "应考", "已考", "未考", "平均分", "通过率(%)"],
            [
                [
                    d.get("rank", ""),
                    d.get("dept_name", ""),
                    d.get("candidate_count", 0),
                    d.get("completed_count", 0),
                    d.get("not_started_count", 0),
                    d.get("average_score", 0),
                    d.get("pass_rate", 0),
                ]
                for d in dept_stats
            ],
        )

    # Sheet 3：按个人（合规留痕）
    if person_stats:
        _write_sheet(
            wb.create_sheet(),
            "按个人",
            ["排名", "姓名", "部门", "得分", "是否通过", "答对", "答错", "用时(秒)", "提交时间", "提交方式", "IP"],
            [
                [
                    p.get("rank", ""),
                    p.get("name", ""),
                    p.get("dept_name", ""),
                    p.get("score", 0),
                    "是" if p.get("is_passed") else "否",
                    p.get("correct_count", 0),
                    p.get("wrong_count", 0),
                    p.get("duration_seconds", 0),
                    str(p.get("submitted_at", "")).replace("T", " ")[:19],
                    p.get("submit_type") or "",
                    p.get("ip_address") or "",
                ]
                for p in person_stats
            ],
        )

    # Sheet 4：按章节
    if chapter_stats:
        _write_sheet(
            wb.create_sheet(),
            "按章节",
            ["章节", "节", "作答次数", "正确次数", "正确率(%)"],
            [
                [
                    c.get("chapter_no", ""),
                    c.get("section_no", ""),
                    c.get("total_answers", 0),
                    c.get("correct_answers", 0),
                    c.get("correct_rate", 0),
                ]
                for c in chapter_stats
            ],
        )

    # Sheet 5：按题目
    if question_stats:
        _write_sheet(
            wb.create_sheet(),
            "按题目",
            ["题号", "题型", "题目", "来源章节", "作答次数", "正确率(%)", "错误率(%)", "需复查"],
            [
                [
                    q.get("order", ""),
                    _type_label(q.get("question_type", "")),
                    q.get("content", ""),
                    f"{q.get('chapter_no') or ''} {q.get('clause_no') or ''}".strip(),
                    q.get("total_answers", 0),
                    q.get("correct_rate", 0),
                    q.get("error_rate", 0),
                    "是" if q.get("needs_review") else "否",
                ]
                for q in question_stats
            ],
        )

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()